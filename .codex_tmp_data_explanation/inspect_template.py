from __future__ import annotations

import hashlib
import json
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


REFERENCE = Path(r"C:\Users\MAHATVA GOEL\.codex\plugins\cache\openai-curated-remote\openai-templates\0.1.1\skills\artifact-template-experiment-analysis\assets\reference.docx")
OUTPUT = Path(r"C:\Users\MAHATVA GOEL\Desktop\Gait Analysis\.codex_tmp_data_explanation\template_structure.json")


def length_inches(value):
    return None if value is None else round(value.inches, 4)


def color_value(font):
    try:
        return None if font.color.rgb is None else str(font.color.rgb)
    except Exception:
        return None


def style_info(style):
    paragraph_format = getattr(style, "paragraph_format", None)
    font = getattr(style, "font", None)
    return {
        "name": style.name,
        "type": str(style.type),
        "base_style": style.base_style.name if style.base_style else None,
        "font": None if font is None else {
            "name": font.name,
            "size_pt": None if font.size is None else round(font.size.pt, 2),
            "bold": font.bold,
            "italic": font.italic,
            "color": color_value(font),
        },
        "paragraph": None if paragraph_format is None else {
            "alignment": None if paragraph_format.alignment is None else str(paragraph_format.alignment),
            "space_before_pt": None if paragraph_format.space_before is None else round(paragraph_format.space_before.pt, 2),
            "space_after_pt": None if paragraph_format.space_after is None else round(paragraph_format.space_after.pt, 2),
            "line_spacing": paragraph_format.line_spacing,
            "keep_with_next": paragraph_format.keep_with_next,
            "keep_together": paragraph_format.keep_together,
            "page_break_before": paragraph_format.page_break_before,
            "left_indent_in": length_inches(paragraph_format.left_indent),
            "right_indent_in": length_inches(paragraph_format.right_indent),
            "first_line_indent_in": length_inches(paragraph_format.first_line_indent),
        },
    }


def para_info(index, para):
    return {
        "index": index,
        "style": para.style.name if para.style else None,
        "text": para.text,
        "alignment": None if para.alignment is None else str(para.alignment),
        "runs": [
            {
                "text": run.text,
                "font": run.font.name,
                "size_pt": None if run.font.size is None else round(run.font.size.pt, 2),
                "bold": run.bold,
                "italic": run.italic,
                "color": color_value(run.font),
            }
            for run in para.runs
        ],
    }


def table_info(index, table):
    grid = table._tbl.tblGrid
    grid_cols = []
    if grid is not None:
        for col in grid.gridCol_lst:
            value = col.get(qn("w:w"))
            grid_cols.append(value)
    rows = []
    for row_index, row in enumerate(table.rows):
        rows.append({
            "row": row_index,
            "cells": [
                {
                    "text": cell.text,
                    "width_in": length_inches(cell.width),
                    "vertical_alignment": None if cell.vertical_alignment is None else str(cell.vertical_alignment),
                }
                for cell in row.cells
            ],
        })
    return {
        "index": index,
        "style": table.style.name if table.style else None,
        "rows": len(table.rows),
        "cols": len(table.columns),
        "grid_widths_twips": grid_cols,
        "content": rows,
    }


def main():
    document = Document(REFERENCE)
    style_counts = Counter(p.style.name if p.style else "" for p in document.paragraphs)
    used_styles = set(style_counts)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    used_styles.add(para.style.name if para.style else "")

    sections = []
    for index, section in enumerate(document.sections):
        sections.append({
            "index": index,
            "page_width_in": length_inches(section.page_width),
            "page_height_in": length_inches(section.page_height),
            "left_margin_in": length_inches(section.left_margin),
            "right_margin_in": length_inches(section.right_margin),
            "top_margin_in": length_inches(section.top_margin),
            "bottom_margin_in": length_inches(section.bottom_margin),
            "header_distance_in": length_inches(section.header_distance),
            "footer_distance_in": length_inches(section.footer_distance),
            "different_first_page_header_footer": section.different_first_page_header_footer,
            "header_paragraphs": [para_info(i, p) for i, p in enumerate(section.header.paragraphs)],
            "footer_paragraphs": [para_info(i, p) for i, p in enumerate(section.footer.paragraphs)],
            "footer_tables": [table_info(i, t) for i, t in enumerate(section.footer.tables)],
        })

    parts = []
    with zipfile.ZipFile(REFERENCE) as archive:
        for info in sorted(archive.infolist(), key=lambda item: item.filename):
            data = archive.read(info.filename)
            parts.append({
                "path": info.filename,
                "size": len(data),
                "sha256": hashlib.sha256(data).hexdigest(),
            })

    payload = {
        "reference": str(REFERENCE),
        "reference_sha256": hashlib.sha256(REFERENCE.read_bytes()).hexdigest(),
        "paragraph_count": len(document.paragraphs),
        "paragraph_style_counts": dict(style_counts),
        "paragraphs": [para_info(i, p) for i, p in enumerate(document.paragraphs)],
        "tables": [table_info(i, t) for i, t in enumerate(document.tables)],
        "sections": sections,
        "styles": [style_info(document.styles[name]) for name in sorted(used_styles) if name in document.styles],
        "package_parts": parts,
    }
    OUTPUT.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    print(json.dumps({
        "paragraph_count": payload["paragraph_count"],
        "table_count": len(payload["tables"]),
        "section_count": len(payload["sections"]),
        "package_part_count": len(parts),
        "used_styles": sorted(used_styles),
    }, indent=2))


if __name__ == "__main__":
    main()
